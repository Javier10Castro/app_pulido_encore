"""Exportación de la base de datos a un libro de Excel con tres hojas: Reporte, Empleados y Materiales."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import database as db

HEAD_FILL = PatternFill("solid", fgColor="1D5BD6")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
BODY_FONT = Font(size=10)
ZEBRA = PatternFill("solid", fgColor="EAF0FB")


def _fill(ws, headers, rows):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY_FONT
            if r % 2 == 0:
                cell.fill = ZEBRA
    for c in range(1, len(headers) + 1):
        vals = [headers[c - 1]] + [str(r[c - 1]) for r in rows]
        ws.column_dimensions[get_column_letter(c)].width = min(34, max(map(len, vals), default=8) + 2)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def export_xlsx(path):
    wb = Workbook()
    _fill(wb.active, ["ID", "Empleado", "Nombre", "Turno", "Material", "Cantidad", "Fecha"],
          [[r["id"], r["empleado"], r["nombre"], r["turno"], r["material"], r["cantidad"], r["fecha"]]
           for r in db.report_rows()])
    wb.active.title = "Reporte"
    _fill(wb.create_sheet("Empleados"), ["ID", "Nombre", "Turno"],
          [[e["id"], e["nombre"], e["turno"]] for e in db.list_employees()])
    _fill(wb.create_sheet("Materiales"), ["Material"],
          [[m] for m in db.list_materials()])
    wb.save(path)
    return path